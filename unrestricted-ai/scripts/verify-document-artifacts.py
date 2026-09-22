"""Verify document creation through the real sandbox and downloadable artifact bytes."""
import hashlib
import json
from pathlib import Path
import time
from urllib.parse import urlsplit
import requests

code='''
from pathlib import Path
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from reportlab.pdfgen import canvas
from pypdf import PdfReader

root=Path('/workspace/document-capability-check')
root.mkdir(exist_ok=True)
marker='DOCUMENT-ARTIFACT-5739'
doc=Document(); doc.add_paragraph(marker); doc.save(root/'check.docx')
assert Document(root/'check.docx').paragraphs[0].text==marker
book=Workbook(); book.active['A1']=marker; book.save(root/'check.xlsx')
assert load_workbook(root/'check.xlsx').active['A1'].value==marker
slides=Presentation(); slide=slides.slides.add_slide(slides.slide_layouts[0])
slide.shapes.title.text=marker; slides.save(root/'check.pptx')
assert Presentation(root/'check.pptx').slides[0].shapes.title.text==marker
pdf=canvas.Canvas(str(root/'check.pdf')); pdf.drawString(72,720,marker); pdf.save()
assert marker in PdfReader(root/'check.pdf').pages[0].extract_text()
print('DOCUMENT_ROUNDTRIPS_PASSED')
'''
report={'started':time.time(),'passed':False,'files':[]}
try:
    r=requests.post('http://workspace-api:8000/code/execute',json={'code':code},timeout=180)
    r.raise_for_status();result=r.json();report['execution']=result
    assert 'DOCUMENT_ROUNDTRIPS_PASSED' in result.get('stdout','') and not result.get('stderr'),result
    for ext in ('docx','xlsx','pptx','pdf'):
        relative='document-capability-check/check.'+ext
        r=requests.post('http://workspace-api:8000/files/artifact',json={'path':relative},timeout=15)
        r.raise_for_status(); artifact=r.json()
        url=urlsplit(artifact['download_url'])
        assert url.netloc=='localhost:8001'
        download=requests.get('http://workspace-api:8000'+url.path+'?'+url.query,timeout=15)
        download.raise_for_status()
        actual=Path('/workspace',relative).read_bytes()
        assert download.content==actual and len(actual)==artifact['bytes']
        assert 'attachment' in download.headers['Content-Disposition']
        report['files'].append({**artifact,'sha256':hashlib.sha256(actual).hexdigest()})
    for invalid in ('../scripts/environment.ps1','/etc/hostname'):
        r=requests.get('http://workspace-api:8000/files/download',params={'path':invalid},timeout=15)
        assert r.status_code==404, 'Download escaped workspace'
    report['passed']=True
finally:
    report['finished']=time.time()
    Path('/state/document-artifact-verification.json').write_text(json.dumps(report,indent=2))
print(json.dumps(report))
